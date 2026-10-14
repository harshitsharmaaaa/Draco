"""Artifact storage (local filesystem under ARTIFACT_DIR) and format renderers.

Every artifact is written atomically, size-limited, read back and hashed before a reference is
returned; downloads are authorized by owner and expiry. Paths are never exposed to the model.
"""
import csv
import hashlib
import io
import os
import re
import tempfile
import threading
import time
from pathlib import Path
from xml.sax.saxutils import escape

MIME = {"docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "pdf": "application/pdf",
        "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "csv": "text/csv", "png": "image/png"}


class ArtifactError(Exception):
    def __init__(self, message, code="internal"):
        super().__init__(message)
        self.code = code


class ArtifactStore:
    def __init__(self, directory, store, *, max_bytes: int, ttl_days: int):
        self.dir, self.store, self.max_bytes, self.ttl_days = Path(directory), store, max_bytes, ttl_days

    def _safe_path(self, key: str) -> Path:
        base = self.dir.resolve()
        p = (base / key).resolve()
        if base not in p.parents:
            raise ArtifactError("invalid storage key")
        return p

    def save(self, *, user_id, kind, fmt, title, data: bytes, meta=None) -> dict:
        if fmt not in MIME:
            raise ArtifactError(f"unsupported artifact format: {fmt!r}", "invalid_input")
        if not data:
            raise ArtifactError("generator produced no data")
        if len(data) > self.max_bytes:
            raise ArtifactError(f"artifact exceeds the {self.max_bytes}-byte limit", "limit_exceeded")
        owner = hashlib.sha256(user_id.encode()).hexdigest()[:16]
        slug = re.sub(r"[^A-Za-z0-9._-]+", "_", title).strip("_")[:60] or kind
        key = f"{owner}/{os.urandom(8).hex()}_{slug}.{fmt}"
        path = self._safe_path(key)
        path.parent.mkdir(parents=True, exist_ok=True)
        fd, tmp = tempfile.mkstemp(dir=path.parent, suffix=".tmp")
        try:
            with os.fdopen(fd, "wb") as fh:
                fh.write(data)
                fh.flush()
                os.fsync(fh.fileno())
            os.replace(tmp, path)
        finally:
            if os.path.exists(tmp):
                os.unlink(tmp)
        sha = hashlib.sha256(data).hexdigest()
        back = path.read_bytes()  # verify what is on disk is what we generated
        if len(back) != len(data) or hashlib.sha256(back).hexdigest() != sha:
            path.unlink(missing_ok=True)
            raise ArtifactError("artifact verification failed after write")
        return self.store.add_artifact(user_id=user_id, kind=kind, fmt=fmt, filename=f"{slug}.{fmt}", storage_key=key,
                                       size_bytes=len(data), sha256=sha, ttl_days=self.ttl_days, meta=meta)

    def resolve(self, artifact_id: str, user_id: str) -> tuple[dict, Path]:
        rec = self.store.get_artifact(artifact_id, user_id)
        if rec is None or rec["status"] != "ready":
            raise ArtifactError("artifact not found", "not_found")
        if rec["expires_at"] < time.time():
            raise ArtifactError("artifact expired", "not_found")
        path = self._safe_path(rec["storage_key"])
        if not path.is_file() or path.stat().st_size != rec["size_bytes"]:
            raise ArtifactError("artifact file is missing or changed", "not_found")
        return rec, path

    def read_bytes(self, artifact_id, user_id) -> bytes:
        return self.resolve(artifact_id, user_id)[1].read_bytes()

    def cleanup_expired(self) -> int:
        n = 0
        for r in self.store.expired_artifacts():
            try:
                self._safe_path(r["storage_key"]).unlink(missing_ok=True)
            except ArtifactError:
                continue
            self.store.mark_artifact(r["artifact_id"], "deleted")
            n += 1
        return n


def ref_view(rec: dict) -> dict:
    return {"artifact_id": rec["artifact_id"], "kind": rec["kind"], "format": rec["fmt"], "filename": rec["filename"],
            "size_bytes": rec["size_bytes"], "sha256": rec["sha256"], "status": rec["status"],
            "download_path": f"/agent/artifacts/{rec['artifact_id']}", "expires_at": rec["expires_at"]}


# ---- shared helpers -----------------------------------------------------------------------
def reference_lines(items) -> list[str]:
    out = []
    for it in items:
        if it.kind == "private":
            pg = f", p.{it.page_start}" if it.page_start is not None else ""
            out.append(f"[{it.ref}] {it.filename or 'document'} (document_id {it.document_id}, chunk_id {it.chunk_id}{pg})")
        else:
            out.append(f"[{it.ref}] {it.title or it.url} - {it.url} (passage {it.passage_id}, web source)")
    return out


_font_lock = threading.Lock()
_FONT = None


def _pdf_font() -> str:
    """DejaVu Sans when present (wider glyph coverage than Helvetica), else Helvetica."""
    global _FONT
    with _font_lock:
        if _FONT is None:
            _FONT = "Helvetica"
            for p in ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/dejavu/DejaVuSans.ttf"):
                if os.path.exists(p):
                    from reportlab.pdfbase import pdfmetrics
                    from reportlab.pdfbase.ttfonts import TTFont
                    pdfmetrics.registerFont(TTFont("DejaVuSans", p))
                    _FONT = "DejaVuSans"
                    break
    return _FONT


def unrenderable_ratio(text: str) -> float:
    """Share of characters the PDF font cannot draw (used before turning web text into a PDF)."""
    font = _pdf_font()
    if font == "Helvetica":
        bad = sum(1 for ch in text if not ch.isspace() and ord(ch) > 255)
    else:
        from reportlab.pdfbase import pdfmetrics
        cmap = pdfmetrics.getFont(font).face.charToGlyph
        bad = sum(1 for ch in text if not ch.isspace() and ord(ch) not in cmap)
    return bad / max(1, len(text))


# ---- DOCX / PDF ---------------------------------------------------------------------------
def render_docx(title, sections, refs) -> bytes:
    from docx import Document
    d = Document()
    d.add_heading(title, level=0)
    for s in sections:
        d.add_heading(s["heading"], level=min(3, max(1, s.get("level", 1))))
        for p in s.get("paragraphs", []):
            d.add_paragraph(p)
        for b in s.get("bullets", []):
            d.add_paragraph(b, style="List Bullet")
        t = s.get("table")
        if t:
            if t.get("caption"):
                d.add_paragraph(t["caption"]).italic = True
            tbl = d.add_table(rows=1, cols=len(t["columns"]))
            tbl.style = "Table Grid"
            for i, c in enumerate(t["columns"]):
                tbl.rows[0].cells[i].text = str(c)
            for row in t["rows"]:
                cells = tbl.add_row().cells
                for i, v in enumerate(row):
                    cells[i].text = "" if v is None else str(v)
        if s.get("source_refs"):
            para = d.add_paragraph("Sources: " + ", ".join(f"[{r}]" for r in s["source_refs"]))
            para.runs[0].italic = True
    if refs:
        d.add_heading("References", level=1)
        for line in refs:
            d.add_paragraph(line)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def render_pdf(title, sections, refs) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    font = _pdf_font()
    ss = getSampleStyleSheet()
    base = ParagraphStyle("b", parent=ss["BodyText"], fontName=font, fontSize=10, leading=14)
    h = {n: ParagraphStyle(f"h{n}", parent=ss[f"Heading{n}"], fontName=font) for n in (1, 2, 3)}
    story = [Paragraph(escape(title), ParagraphStyle("t", parent=ss["Title"], fontName=font)), Spacer(1, 8)]
    for s in sections:
        story.append(Paragraph(escape(s["heading"]), h[min(3, max(1, s.get("level", 1)))]))
        for p in s.get("paragraphs", []):
            story += [Paragraph(escape(p), base), Spacer(1, 4)]
        if s.get("bullets"):
            story.append(ListFlowable([ListItem(Paragraph(escape(b), base)) for b in s["bullets"]], bulletType="bullet"))
        t = s.get("table")
        if t:
            data = [[Paragraph(escape(str(c)), base) for c in t["columns"]]] + \
                   [[Paragraph(escape("" if v is None else str(v)), base) for v in r] for r in t["rows"]]
            tb = Table(data, repeatRows=1)
            tb.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), .5, colors.grey), ("BACKGROUND", (0, 0), (-1, 0), colors.whitesmoke),
                                    ("VALIGN", (0, 0), (-1, -1), "TOP")]))
            story += [Spacer(1, 4), tb, Spacer(1, 6)]
        if s.get("source_refs"):
            story.append(Paragraph("<i>Sources: " + ", ".join(f"[{r}]" for r in s["source_refs"]) + "</i>", base))
    if refs:
        story.append(Paragraph("References", h[1]))
        story += [Paragraph(escape(line), base) for line in refs]
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, title=title).build(story)
    return buf.getvalue()


def text_to_pdf(title: str, header_lines: list[str], body: str) -> bytes:
    """Used to feed a web snapshot through the existing PDF-only ingestion pipeline."""
    paras = [p.strip() for p in re.split(r"\n\s*\n", body) if p.strip()]
    return render_pdf(title, [{"heading": "Source", "level": 2, "paragraphs": header_lines},
                              {"heading": "Content", "level": 2, "paragraphs": paras}], [])


# ---- PPTX ---------------------------------------------------------------------------------
def render_pptx(title, slides, refs, image_bytes_by_id) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    s0 = prs.slides.add_slide(prs.slide_layouts[0])
    s0.shapes.title.text = title
    if len(s0.placeholders) > 1:
        s0.placeholders[1].text = time.strftime("%Y-%m-%d")
    for sl in slides:
        has_table = bool(sl.get("table"))
        slide = prs.slides.add_slide(prs.slide_layouts[5 if has_table or sl.get("chart_artifact_id") else 1])
        slide.shapes.title.text = sl["title"]
        top = Inches(1.6)
        if not has_table and not sl.get("chart_artifact_id"):
            body = slide.placeholders[1].text_frame
            body.clear()
            for i, b in enumerate(sl.get("bullets", [])):
                p = body.paragraphs[0] if i == 0 else body.add_paragraph()
                p.text, p.font.size = b, Pt(22)
        elif has_table:
            t = sl["table"]
            gs = slide.shapes.add_table(len(t["rows"]) + 1, len(t["columns"]), Inches(.7), top, Inches(11.9),
                                        Inches(.4) * (len(t["rows"]) + 1))
            for j, c in enumerate(t["columns"]):
                gs.table.cell(0, j).text = str(c)
            for i, r in enumerate(t["rows"], 1):
                for j, v in enumerate(r):
                    gs.table.cell(i, j).text = "" if v is None else str(v)
        else:
            img = image_bytes_by_id.get(sl["chart_artifact_id"])
            if img:
                slide.shapes.add_picture(io.BytesIO(img), Inches(2), top, height=Inches(5))
        if sl.get("source_refs"):
            tb = slide.shapes.add_textbox(Inches(.7), Inches(6.9), Inches(11.9), Inches(.4))
            tb.text_frame.text = "Sources: " + ", ".join(f"[{r}]" for r in sl["source_refs"])
            tb.text_frame.paragraphs[0].font.size = Pt(11)
        if sl.get("notes"):
            slide.notes_slide.notes_text_frame.text = sl["notes"]
    if refs:
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = "Sources"
        tf = s.placeholders[1].text_frame
        tf.clear()
        for i, line in enumerate(refs[:14]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text, p.font.size = line, Pt(12)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# ---- XLSX / CSV ---------------------------------------------------------------------------
_FORMULA_OK = re.compile(r"^=[A-Z0-9(),.:+\-*/<>= $%]+$")
_FUNCS_OK = {"SUM", "AVERAGE", "MIN", "MAX", "COUNT", "COUNTA", "ROUND", "ABS", "IF"}
_CELL = re.compile(r"^[A-Z]{1,3}[1-9]\d{0,5}$")


def formula_problem(cell: str, formula: str) -> str | None:
    if not _CELL.match(cell):
        return f"invalid cell address '{cell}'"
    if not _FORMULA_OK.match(formula) or "!" in formula:
        return "formula has characters outside the allowed set (no sheet/external references or strings)"
    bad = {f for f in re.findall(r"([A-Z][A-Z0-9]*)\(", formula)} - _FUNCS_OK
    if bad:
        return f"function(s) not allowed: {sorted(bad)}"
    if formula.count("(") != formula.count(")"):
        return "unbalanced parentheses"
    return None


def _text_safe(v):
    return "'" + v if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r") else v


def render_csv(columns, rows) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([_text_safe(c) for c in columns])
    for r in rows:
        w.writerow(["" if v is None else _text_safe(v) for v in r])
    return buf.getvalue().encode("utf-8-sig")


def render_xlsx(sheets, refs) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    wb.remove(wb.active)
    for sh in sheets:
        ws = wb.create_sheet(sh["name"])
        ws.append(list(sh["columns"]))
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in sh["rows"]:
            ws.append(["" if v is None else v for v in r])
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value[:1] in ("=", "+", "-", "@"):
                    c.data_type = "s"  # literal text, never a formula
        for f in sh.get("formulas", []):
            ws[f["cell"]] = f["formula"]  # validated by formula_problem()
        ws.freeze_panes = "A2"
    if refs:
        ws = wb.create_sheet("Provenance")
        ws.append(["Reference"])
        for line in refs:
            ws.append([line])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---- charts -------------------------------------------------------------------------------
def render_chart(chart_type, title, x_label, y_label, unit, categories, series) -> bytes:
    from matplotlib.backends.backend_agg import FigureCanvasAgg
    from matplotlib.figure import Figure
    fig = Figure(figsize=(8, 4.8), dpi=150)
    FigureCanvasAgg(fig)
    ax = fig.add_subplot(111)
    ylab = f"{y_label} ({unit})" if unit else y_label
    if chart_type == "pie":
        ax.pie(series[0]["values"], labels=categories, autopct="%1.1f%%")
        ax.axis("equal")
    elif chart_type == "scatter":
        for s in series:
            ax.scatter([float(c) for c in categories], s["values"], label=s["name"])
    else:
        n = len(series)
        for i, s in enumerate(series):
            if chart_type == "line":
                ax.plot(categories, s["values"], marker="o", label=s["name"])
            elif chart_type == "hbar":
                ax.barh([j + i * .8 / n for j in range(len(categories))], s["values"], height=.8 / n, label=s["name"])
                ax.set_yticks([j + .4 - .4 / n for j in range(len(categories))], categories)
            else:
                ax.bar([j + i * .8 / n for j in range(len(categories))], s["values"], width=.8 / n, label=s["name"])
                ax.set_xticks([j + .4 - .4 / n for j in range(len(categories))], categories, rotation=30 if len(categories) > 6 else 0)
    if chart_type != "pie":
        (ax.set_ylabel if chart_type != "hbar" else ax.set_xlabel)(ylab)
        (ax.set_xlabel if chart_type != "hbar" else ax.set_ylabel)(x_label)
        if len(series) > 1:
            ax.legend()
        ax.grid(alpha=.3)
    ax.set_title(title)
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png")
    return buf.getvalue()
